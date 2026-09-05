"""Portable preparation exports, including file bytes for remote MCP clients."""

import base64
import json
import tempfile
from pathlib import Path
from typing import Any, Literal


def export_preparation_report(
    report: dict[str, Any], format_name: Literal["markdown", "json", "excel", "pdf"]
) -> dict[str, str]:
    extension, mime = {
        "markdown": ("md", "text/markdown"), "json": ("json", "application/json"),
        "excel": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "pdf": ("pdf", "application/pdf"),
    }[format_name]
    # A unique directory prevents collisions and excludes user-controlled paths.
    path = Path(tempfile.mkdtemp(prefix="vgc-preparation-")) / f"team-preparation.{extension}"
    if format_name in ("markdown", "json"):
        content = report["report_markdown"] if format_name == "markdown" else json.dumps(report, indent=2)
        path.write_text(content, encoding="utf-8")
    elif format_name == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Preparation"
        # Explicit string cell type prevents formula execution from pasted names.
        for index, line in enumerate(report["report_markdown"].splitlines(), 1):
            cell = sheet.cell(index, 1, line)
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if line.startswith("#"):
                cell.font = Font(bold=True)
            sheet.row_dimensions[index].height = max(18, 15 * (1 + len(line) // 110))
        sheet.column_dimensions["A"].width = 125
        sheet.freeze_panes = "A2"
        matchups = workbook.create_sheet("Matchups")
        matchups.append(["Pokemon", "Opponent", "Outgoing move", "Damage dealt", "OHKO %", "Incoming move", "Damage taken", "Incoming OHKO %", "Raw Speed faster"])
        for row in report["matchups"]:
            outgoing, incoming = row["outgoing"] or {}, row["incoming"] or {}
            matchups.append([row["pokemon"], row["opponent"], outgoing.get("move"), outgoing.get("damage"),
                             outgoing.get("ohko_percent"), incoming.get("move"), incoming.get("damage"),
                             incoming.get("ohko_percent"), row["faster"]])
        matchups.auto_filter.ref = matchups.dimensions
        matchups.freeze_panes = "C2"
        for row in matchups:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="233B55")
                elif cell.column in (5, 8) and isinstance(cell.value, (float, int)):
                    color = "C6EFCE" if cell.value == 100 else "FFEB9C" if cell.value > 0 else "F2F2F2"
                    cell.fill = PatternFill("solid", fgColor=color)
        for index in range(1, 10):
            matchups.column_dimensions[get_column_letter(index)].width = 28 if index in (4, 7) else 22
        matchups.sheet_properties.pageSetUpPr.fitToPage = True
        matchups.page_setup.orientation = "landscape"
        matchups.page_setup.fitToWidth = 1
        workbook.save(path)
    else:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        def paragraph(text: str, *, heading: bool = False, code: bool = False) -> None:
            pdf.set_font("Courier" if code else "Helvetica", style="B" if heading else "", size=12 if heading else 9)
            printable = text.encode("latin-1", errors="replace").decode("latin-1")
            pdf.multi_cell(0, 5, printable or " ", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        paragraph(f"Team preparation - {report['regulation']}", heading=True)
        source = report["source"]
        paragraph(f"Chaos usage: {source['month']} | {source['format']} | rating {source['rating']}")
        paragraph(source["source_url"])
        paragraph(report["legality"]["message"])
        paragraph("Your team", heading=True)
        paragraph(report["showdown_paste"], code=True)
        paragraph("Copy this and paste it directly into Pokemon Showdown's teambuilder.")
        paragraph("Threat comparisons", heading=True)
        for row in report["matchups"]:
            paragraph(f"{row['pokemon']} vs {row['opponent']}", heading=True)
            for label, value in (("Dealt", row["outgoing"]), ("Taken", row["incoming"])):
                paragraph(f"{label}: {value['move']} - {value['damage']} ({value['ohko_percent']:.2f}% OHKO)" if value else f"{label}: no damaging moves")
        paragraph("Exact opponent spreads", heading=True)
        seen = set()
        for row in report["matchups"]:
            if row["opponent_showdown_paste"] not in seen:
                paragraph(row["opponent_showdown_paste"], code=True)
                seen.add(row["opponent_showdown_paste"])
        for adjustment in report["verified_adjustments"]:
            paragraph(f"Verified alternatives - {adjustment['pokemon']}", heading=True)
            for candidate in adjustment.get("candidates", []):
                paragraph(candidate["purpose"])
                paragraph(candidate["showdown_paste"], code=True)
        paragraph("Assumptions and validation scope", heading=True)
        for line in [report["legality"]["scope"], *report["assumptions"], *report["legality"]["violations"], *report["legality"]["warnings"]]:
            paragraph(line)
        pdf.output(str(path))
    return {"filename": path.name, "mime_type": mime, "file_path": str(path),
            "data_base64": base64.b64encode(path.read_bytes()).decode("ascii"),
            "delivery": "Decode data_base64 to create a downloadable file in the MCP client."}
